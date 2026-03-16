from flask import Blueprint, render_template, request, send_file
from flask_login import login_required
from models import db, Vehicle, Sale, Quotation, Customer, Payment, CostSheet
from sqlalchemy import func, extract
from datetime import date, timedelta, datetime
import io

reports_bp = Blueprint('reports', __name__)


def get_date_range():
    start_str = request.args.get('start')
    end_str = request.args.get('end')
    if start_str:
        start = datetime.strptime(start_str, '%Y-%m-%d')
    else:
        start = datetime(date.today().year, 1, 1)
    if end_str:
        end = datetime.strptime(end_str, '%Y-%m-%d')
    else:
        end = datetime.now()
    return start, end, start_str or start.date().isoformat(), end_str or end.date().isoformat()


@reports_bp.route('/')
@login_required
def dashboard():
    today = date.today()
    month_start = date(today.year, today.month, 1)

    # KPIs
    total_inventory = Vehicle.query.filter_by(is_archived=False).filter(Vehicle.status != 'sold').count()
    active_listings = Vehicle.query.filter_by(is_archived=False, status='available').count()
    monthly_sales = Sale.query.filter(Sale.status.in_(['full_payment', 'delivered']),
                                      Sale.created_at >= month_start).count()
    monthly_revenue = db.session.query(func.sum(Sale.sale_price)).filter(
        Sale.status.in_(['full_payment', 'delivered']),
        Sale.created_at >= month_start
    ).scalar() or 0

    total_customers = Customer.query.count()
    pending_quotes = Quotation.query.filter(Quotation.status.in_(['draft', 'sent'])).count()
    overdue_quotes = Quotation.query.filter(
        Quotation.valid_until < today,
        Quotation.status.in_(['draft', 'sent'])
    ).count()

    # Pipeline counts — only active (non-terminal) stages
    active_stages = ['lead', 'negotiation', 'deposit_paid', 'full_payment']
    pipeline_data = {}
    for stage in active_stages:
        pipeline_data[stage] = Sale.query.filter_by(status=stage).count()

    # Monthly sales chart (last 6 months)
    monthly_chart = []
    for i in range(5, -1, -1):
        d = today.replace(day=1) - timedelta(days=i * 28)
        m_start = date(d.year, d.month, 1)
        if d.month == 12:
            m_end = date(d.year + 1, 1, 1)
        else:
            m_end = date(d.year, d.month + 1, 1)
        rev = db.session.query(func.sum(Sale.sale_price)).filter(
            Sale.status.in_(['full_payment', 'delivered']),
            Sale.created_at >= m_start, Sale.created_at < m_end
        ).scalar() or 0
        cnt = Sale.query.filter(
            Sale.status.in_(['full_payment', 'delivered']),
            Sale.created_at >= m_start, Sale.created_at < m_end
        ).count()
        monthly_chart.append({'month': m_start.strftime('%b %Y'), 'revenue': rev, 'count': cnt})

    # Inventory by status
    inv_status = {
        'available': Vehicle.query.filter_by(is_archived=False, status='available').count(),
        'reserved': Vehicle.query.filter_by(is_archived=False, status='reserved').count(),
        'under_inspection': Vehicle.query.filter_by(is_archived=False, status='under_inspection').count(),
        'sold': Vehicle.query.filter_by(status='sold').count(),
    }

    # Aging inventory (days since added)
    aging_30 = Vehicle.query.filter_by(is_archived=False, status='available').filter(
        Vehicle.created_at <= datetime.now() - timedelta(days=30)).count()
    aging_60 = Vehicle.query.filter_by(is_archived=False, status='available').filter(
        Vehicle.created_at <= datetime.now() - timedelta(days=60)).count()
    aging_90 = Vehicle.query.filter_by(is_archived=False, status='available').filter(
        Vehicle.created_at <= datetime.now() - timedelta(days=90)).count()

    # Top makes
    top_makes = db.session.query(Vehicle.make, func.count(Vehicle.id).label('cnt')
        ).filter_by(is_archived=False).group_by(Vehicle.make).order_by(func.count(Vehicle.id).desc()).limit(5).all()

    # Recent sales
    recent_sales = Sale.query.filter(Sale.status.in_(['full_payment', 'delivered'])
                                     ).order_by(Sale.created_at.desc()).limit(5).all()

    # Serialisable versions for Chart.js
    pipeline_chart = {
        'labels': [k.replace('_', ' ').title() for k in pipeline_data.keys()],
        'counts': list(pipeline_data.values()),
    }
    inv_status_chart = {
        'labels': [k.replace('_', ' ').title() for k in inv_status.keys()],
        'counts': list(inv_status.values()),
    }
    monthly_chart_js = {
        'labels': [m['month'] for m in monthly_chart],
        'revenue': [m['revenue'] for m in monthly_chart],
        'units': [m['count'] for m in monthly_chart],
    }

    stats = {
        'total_inventory': total_inventory,
        'active_listings': active_listings,
        'monthly_sales': monthly_sales,
        'monthly_revenue': monthly_revenue,
        'total_customers': total_customers,
        'pending_quotes': pending_quotes,
        'overdue_quotes': overdue_quotes,
        'pipeline_total': sum(pipeline_data.values()),  # active deals only
        'pipeline_data': pipeline_data,
        'aging_30': aging_30,
        'aging_60': aging_60,
        'aging_90': aging_90,
    }
    return render_template('admin/reports/dashboard.html',
                           stats=stats,
                           monthly_chart=monthly_chart_js,
                           pipeline_chart=pipeline_chart,
                           inv_status_chart=inv_status_chart,
                           inv_status=inv_status,
                           top_makes=top_makes, recent_sales=recent_sales)


@reports_bp.route('/sales')
@login_required
def sales_report():
    start, end, start_str, end_str = get_date_range()
    sales = Sale.query.filter(
        Sale.created_at >= start, Sale.created_at <= end
    ).order_by(Sale.created_at.desc()).all()

    total_revenue = sum(s.sale_price or 0 for s in sales if s.status in ['full_payment', 'delivered'])
    completed = [s for s in sales if s.status in ['full_payment', 'delivered']]

    # By salesperson
    by_person = {}
    for s in completed:
        name = s.salesperson.name if s.salesperson else 'Unassigned'
        if name not in by_person:
            by_person[name] = {'count': 0, 'revenue': 0}
        by_person[name]['count'] += 1
        by_person[name]['revenue'] += s.sale_price or 0

    # By payment method
    by_method = {}
    for s in completed:
        for p in s.payments:
            m = p.payment_method or 'Unknown'
            by_method[m] = by_method.get(m, 0) + p.amount

    total_collected = sum(
        p.amount for s in completed for p in s.payments
    )
    summary = type('S', (), {
        'total_sales': len(completed),
        'total_revenue': total_revenue,
        'avg_sale_price': total_revenue / len(completed) if completed else 0,
        'total_collected': total_collected,
    })()

    return render_template('admin/reports/sales.html',
                           sales=sales, start_str=start_str, end_str=end_str,
                           total_revenue=total_revenue, completed=completed,
                           by_person=by_person, by_method=by_method, summary=summary)


@reports_bp.route('/inventory')
@login_required
def inventory_report():
    today = datetime.now()
    vehicles = Vehicle.query.filter_by(is_archived=False, status='available').order_by(Vehicle.created_at).all()
    for v in vehicles:
        v.days_in_stock = (today - v.created_at).days

    aging_data = {
        '0-30 days': [v for v in vehicles if v.days_in_stock <= 30],
        '31-60 days': [v for v in vehicles if 31 <= v.days_in_stock <= 60],
        '61-90 days': [v for v in vehicles if 61 <= v.days_in_stock <= 90],
        '90+ days': [v for v in vehicles if v.days_in_stock > 90],
    }
    stats = type('S', (), {
        'total': Vehicle.query.filter_by(is_archived=False).count(),
        'available': Vehicle.query.filter_by(is_archived=False, status='available').count(),
        'reserved': Vehicle.query.filter_by(is_archived=False, status='reserved').count(),
        'sold': Vehicle.query.filter_by(status='sold').count(),
    })()
    return render_template('admin/reports/inventory.html',
                           vehicles=vehicles, aging_data=aging_data, stats=stats)


@reports_bp.route('/profitability')
@login_required
def profitability_report():
    start, end, start_str, end_str = get_date_range()
    completed_sales = Sale.query.filter(
        Sale.status.in_(['full_payment', 'delivered']),
        Sale.created_at >= start, Sale.created_at <= end
    ).all()

    report_data = []
    total_landed = 0
    total_selling = 0
    for s in completed_sales:
        cs = s.vehicle.cost_sheet if s.vehicle else None
        landed = cs.total_landed_cost if cs else 0
        selling = s.sale_price or 0
        profit = selling - landed
        margin = (profit / selling * 100) if selling > 0 else 0
        report_data.append(type('R', (), {
            'sale_id': s.id,
            'sale_ref': s.sale_ref,
            'vehicle': s.vehicle,
            'landed_cost': landed,
            'sale_price': selling,
            'profit': profit,
            'margin': margin,
        })())
        total_landed += landed
        total_selling += selling

    total_profit = total_selling - total_landed
    avg_margin = (total_profit / total_selling * 100) if total_selling > 0 else 0

    summary = type('S', (), {
        'total_revenue': total_selling,
        'total_cost': total_landed,
        'gross_profit': total_profit,
        'avg_margin': avg_margin,
    })()

    return render_template('admin/reports/profitability.html',
                           rows=report_data, start_str=start_str, end_str=end_str,
                           total_landed=total_landed, total_selling=total_selling,
                           total_profit=total_profit, avg_margin=avg_margin, summary=summary)


@reports_bp.route('/export/sales')
@login_required
def export_sales():
    start, end, start_str, end_str = get_date_range()
    sales = Sale.query.filter(Sale.created_at >= start, Sale.created_at <= end).order_by(Sale.created_at.desc()).all()

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sales Report'
        headers = ['Sale Ref', 'Date', 'Customer', 'Vehicle', 'Salesperson', 'Status', 'Sale Price (LKR)', 'Total Paid (LKR)', 'Balance (LKR)']
        header_fill = PatternFill(start_color='1a3a5c', end_color='1a3a5c', fill_type='solid')
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
        for row, s in enumerate(sales, 2):
            ws.cell(row=row, column=1, value=s.sale_ref)
            ws.cell(row=row, column=2, value=s.created_at.strftime('%d %b %Y'))
            ws.cell(row=row, column=3, value=s.customer.name if s.customer else '')
            ws.cell(row=row, column=4, value=f'{s.vehicle.year} {s.vehicle.make} {s.vehicle.model}' if s.vehicle else '')
            ws.cell(row=row, column=5, value=s.salesperson.name if s.salesperson else '')
            ws.cell(row=row, column=6, value=s.status_label)
            ws.cell(row=row, column=7, value=s.sale_price or 0)
            ws.cell(row=row, column=8, value=s.total_paid())
            ws.cell(row=row, column=9, value=s.balance_due())
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].auto_size = True
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=f'sales_report_{start_str}_{end_str}.xlsx')
    except ImportError:
        return 'openpyxl not installed', 500
